"""Build the status deck from the SLAC template (python-pptx).

    python slides/build_deck.py     (from the repo root; needs python-pptx, figs/ from collect.py + make_figs.py)

The template deck provides three frames: the title slide, the content slide (title, footer, slide number) and the
"Back ups" divider. Content slides are copies of the content frame; every other template slide is dropped.
Numbers come from figs/summary.json (slides/collect.py) unless marked otherwise in the text.
"""
import copy
import json
import os

from lxml import etree
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt

HERE = os.path.dirname(os.path.abspath(__file__))
FIG = os.path.join(HERE, "figs")
TEMPLATE = "/pscratch/sd/l/liangyu/FCC/PFlow/slides/Agentic_opt_CLD_Pandora.pptx"
OUT = os.path.join(HERE, "CVAE_HAXAD_status.pptx")
DATE = "19 Oct 2026"

RED, RED2, DARK, BLACK, GREY, LIGHT = "8C1515", "B83A4B", "3F3F3F", "000000", "BFBFBF", "F2F2F2"
FONT = "Lato"
S = json.load(open(os.path.join(FIG, "summary.json")))
RUN, GRP = S["runs"], S["groups"]


def f3(x):
    return f"{x:.3f}"


def f2(x):
    return f"{x:.2f}".replace("-0.00", "0.00")


def ms(g, m, nd=3):
    s = GRP[g][m]
    return f"{s['mean']:.{nd}f} ± {s['std']:.{nd}f}"


# --- helpers (same look as the Pandora deck) ------------------------------------------------------------------
def rgb(h):
    return RGBColor.from_string(h)


def clear(slide):
    for sh in list(slide.shapes):
        if not sh.is_placeholder:
            sh._element.getparent().remove(sh._element)


def set_title(slide, text):
    tf = slide.shapes.title.text_frame
    p = tf.paragraphs[0]
    for r in list(p.runs)[1:]:
        r._r.getparent().remove(r._r)
    p.runs[0].text = text
    for extra in list(tf.paragraphs)[1:]:
        extra._p.getparent().remove(extra._p)


def _style_run(run, size, color, bold=False, italic=False, font=FONT):
    run.font.size = Pt(size)
    run.font.name = font
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = rgb(color)
    rPr = run._r.get_or_add_rPr()
    for tag in ("a:ea", "a:cs"):
        el = rPr.find(qn(tag))
        if el is None:
            el = etree.SubElement(rPr, qn(tag))
        el.set("typeface", font)


def _bullet(p, level):
    pPr = p._p.get_or_add_pPr()
    marL, indent = (Emu(228600), Emu(-228600)) if level == 0 else (Emu(571500), Emu(-228600))
    pPr.set("marL", str(marL)); pPr.set("indent", str(indent)); pPr.set("algn", "l")
    for child in list(pPr):
        pPr.remove(child)
    ln = etree.SubElement(pPr, qn("a:lnSpc")); etree.SubElement(ln, qn("a:spcPct")).set("val", "110000")
    sb = etree.SubElement(pPr, qn("a:spcBef")); etree.SubElement(sb, qn("a:spcPts")).set("val", "600" if level == 0 else "300")
    bc = etree.SubElement(pPr, qn("a:buClr")); etree.SubElement(bc, qn("a:srgbClr")).set("val", BLACK if level == 0 else RED2)
    etree.SubElement(pPr, qn("a:buSzPct")).set("val", "100000")
    etree.SubElement(pPr, qn("a:buFont")).set("typeface", "Arial")
    etree.SubElement(pPr, qn("a:buChar")).set("char", "•")


def add_bullets(slide, x, y, w, h, items, size=14):
    """items: list of (level, text | [(text, bold), ...])."""
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = Emu(0); tf.margin_top = tf.margin_bottom = Emu(45700)
    for i, (level, content) in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        _bullet(p, level)
        for text, bold in (content if isinstance(content, list) else [(content, False)]):
            r = p.add_run(); r.text = text
            _style_run(r, size if level == 0 else size - 1, BLACK if level == 0 else DARK, bold=bold)
    return box


def add_text(slide, x, y, w, h, text, size=11, color=DARK, bold=False, italic=False, align=PP_ALIGN.LEFT):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame; tf.word_wrap = True
    tf.margin_left = tf.margin_right = Emu(0); tf.margin_top = tf.margin_bottom = Emu(0)
    for i, line in enumerate(text if isinstance(text, list) else [text]):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        r = p.add_run(); r.text = line
        _style_run(r, size, color, bold=bold, italic=italic)
    return box


def add_note(slide, x, y, w, h, text, size=10):
    """The template's yellow call-out box."""
    box = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    box.fill.solid(); box.fill.fore_color.rgb = rgb("FFFF00")
    box.line.color.rgb = rgb("CC0000"); box.line.width = Pt(0.75)
    box.shadow.inherit = False
    tf = box.text_frame; tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = Emu(64000)
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
    r = p.add_run(); r.text = text
    _style_run(r, size, BLACK, font="Comfortaa")
    return box


def add_table(slide, x, y, w, col_w, rows, size=11, row_h=0.3, bold_rows=()):
    n_r, n_c = len(rows), len(rows[0])
    shape = slide.shapes.add_table(n_r, n_c, Inches(x), Inches(y), Inches(w), Inches(row_h * n_r))
    tbl = shape.table
    tblPr = tbl._tbl.tblPr
    tblPr.set("firstRow", "1"); tblPr.set("bandRow", "0")
    sid = tblPr.find(qn("a:tableStyleId"))
    if sid is None:
        sid = etree.SubElement(tblPr, qn("a:tableStyleId"))
    sid.text = "{5940675A-B579-460E-94D1-54222C63F5DA}"
    total = sum(col_w)
    for i, cw in enumerate(col_w):
        tbl.columns[i].width = Inches(w * cw / total)
    for i in range(n_r):
        tbl.rows[i].height = Inches(row_h)
        for j in range(n_c):
            cell = tbl.cell(i, j)
            cell.margin_left = cell.margin_right = Emu(54000); cell.margin_top = cell.margin_bottom = Emu(18000)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            head = i == 0
            cell.fill.solid(); cell.fill.fore_color.rgb = rgb(RED if head else ("FFFFFF" if i % 2 else LIGHT))
            tf = cell.text_frame; tf.word_wrap = True
            r = tf.paragraphs[0].add_run(); r.text = str(rows[i][j])
            _style_run(r, size, "FFFFFF" if head else DARK, bold=head or i in bold_rows)
    return shape


def add_image(slide, name, x, y, w=None, h=None):
    kw = {}
    if w: kw["width"] = Inches(w)
    if h: kw["height"] = Inches(h)
    return slide.shapes.add_picture(os.path.join(FIG, name), Inches(x), Inches(y), **kw)


def add_steps(slide, x, y, w, h, steps, status):
    """Row of chevrons; status: 'done' (red), 'now' (light red), 'todo' (grey)."""
    n = len(steps); gap = 0.05; cw = (w - gap * (n - 1)) / n
    for i, (label, st) in enumerate(zip(steps, status)):
        shp = slide.shapes.add_shape(MSO_SHAPE.PENTAGON if i == 0 else MSO_SHAPE.CHEVRON,
                                     Inches(x + i * (cw + gap)), Inches(y), Inches(cw), Inches(h))
        shp.fill.solid(); shp.fill.fore_color.rgb = rgb({"done": RED, "now": RED2, "todo": GREY}[st])
        shp.line.fill.background(); shp.shadow.inherit = False
        tf = shp.text_frame; tf.word_wrap = True
        tf.margin_left = tf.margin_right = Emu(30000); tf.margin_top = tf.margin_bottom = Emu(0)
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
        r = p.add_run(); r.text = label
        _style_run(r, 10, "FFFFFF" if st != "todo" else DARK, bold=True)


def add_box(slide, x, y, w, h, head, body, fill=LIGHT, line=RED):
    shp = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    shp.fill.solid(); shp.fill.fore_color.rgb = rgb(fill); shp.line.color.rgb = rgb(line); shp.line.width = Pt(1)
    shp.shadow.inherit = False
    tf = shp.text_frame; tf.word_wrap = True; tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    tf.margin_left = tf.margin_right = Emu(50000)
    p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
    r = p.add_run(); r.text = head; _style_run(r, 12, RED, bold=True)
    if body:
        p2 = tf.add_paragraph(); p2.alignment = PP_ALIGN.CENTER
        r = p2.add_run(); r.text = body; _style_run(r, 10, DARK)
    return shp


def add_arrow(slide, x, y, w, h=0.28, kind=MSO_SHAPE.RIGHT_ARROW, color=RED2):
    ar = slide.shapes.add_shape(kind, Inches(x), Inches(y), Inches(w), Inches(h))
    ar.fill.solid(); ar.fill.fore_color.rgb = rgb(color); ar.line.fill.background(); ar.shadow.inherit = False
    return ar


# --- frames ------------------------------------------------------------------------------------------------------
prs = Presentation(TEMPLATE)
T = list(prs.slides)
title_s, frame, backup_s = T[0], T[1], T[17]


def new_content():
    """Copy of the content frame (title, footer, slide number), without the frame's own content."""
    s = prs.slides.add_slide(frame.slide_layout)
    for sh in list(s.shapes):
        sh._element.getparent().remove(sh._element)
    for el in frame.shapes._spTree.iterchildren():
        if el.tag.endswith("}sp") and el.find(".//" + qn("p:ph")) is not None:
            s.shapes._spTree.append(copy.deepcopy(el))
    return s


# --- 1 title -------------------------------------------------------------------------------------------------------
ph = {sh.placeholder_format.idx: sh for sh in title_s.placeholders}
paras = list(ph[1].text_frame.paragraphs)
for p, text in zip(paras[:2], ["A Single ParT CVAE for HAXAD:", "Embedding and Generation"]):
    p.runs[0].text = text
    for r in p.runs[1:]:
        r._r.getparent().remove(r._r)
for p in paras[2:]:
    p._p.getparent().remove(p._p)
for r in paras[1].runs:
    r.font.size = Pt(41)
clear(title_s)
add_text(title_s, 0.88, 3.45, 9, 0.4, "Status: reconstruction, z–mγγ decorrelation, multi-seed baseline", size=16, italic=True)
date_ph = [sh for sh in title_s.placeholders if sh.text_frame.text.strip() == "17 Oct 2026"]
for sh in date_ph:
    sh.text_frame.paragraphs[0].runs[0].text = DATE

slides = []

# --- 2 goal --------------------------------------------------------------------------------------------------------
s = new_content(); slides.append(s); set_title(s, "Goal: One Latent Space for Embedding and Generation")
add_bullets(s, 0.6, 1.15, 7.0, 5.4, [
    (0, [("Today", True), (": HAXAD uses a contrastive embedding model and a separate generative model, conditioned on mγγ, "
                           "to build the signal-region (SR) background template from the sidebands (SB)", False)]),
    (0, [("This work", True), (": a single conditional VAE with latent [z (16 free dims), mγγ (one fixed dim)]", False)]),
    (1, "encoder: event → z, plus one output that reproduces mγγ"),
    (1, "decoder: (z, mγγ) → event, so events can be generated at a specified mγγ"),
    (0, [("Downstream use", True), (": the anomaly classifier runs on z; the SR background template is the SB z distribution", False)]),
    (0, [("Requirements", True), (", in order of priority:", False)]),
    (1, "z independent of mγγ: a classifier on z cannot tell SR from SB (closure AUC → 0.5)"),
    (1, "z keeps the event information (reconstruction R²), hence the signal information"),
    (1, "only mγγ is fixed in the latent; nothing else has to be specified at generation"),
], size=13)
add_box(s, 8.1, 1.4, 4.6, 0.85, "SB events (105–120, 130–145 GeV)", "train the CVAE; z_SB")
add_arrow(s, 10.25, 2.33, 0.3, 0.32, MSO_SHAPE.DOWN_ARROW)
add_box(s, 8.1, 2.75, 4.6, 0.85, "SR template", "z_SB; valid in the SR only if z is independent of mγγ")
add_arrow(s, 10.25, 3.68, 0.3, 0.32, MSO_SHAPE.DOWN_ARROW)
add_box(s, 8.1, 4.1, 4.6, 0.85, "Weakly supervised classifier on z", "SR data vs template → anomaly score")
add_note(s, 8.1, 5.3, 4.6, 0.75, "Closure test: a BDT on z separates SR from SB background. Any AUC above 0.5 is a fake excess in the SR.")

# --- 3 model and data ----------------------------------------------------------------------------------------------
s = new_content(); slides.append(s); set_title(s, "Model and Data")
add_box(s, 0.6, 1.25, 2.7, 1.25, "Inputs", "11 object slots + 2 photon slots,\npair interactions, 7 high-level")
add_arrow(s, 3.38, 1.73, 0.4)
add_box(s, 3.85, 1.25, 2.7, 1.25, "ParT encoder", "DPTModel (contrastive encoder)\n4 self- + 2 class-attention blocks")
add_arrow(s, 6.63, 1.73, 0.4)
add_box(s, 7.1, 1.25, 2.4, 1.25, "Latent", "z: μ, σ (16 dims)\nm̂γγ (regressed, 1 dim)")
add_arrow(s, 9.58, 1.73, 0.4)
add_box(s, 10.05, 1.25, 2.75, 1.25, "Slot-query decoder", "[z, MLP(mγγ)] → 13 slots +\nhigh level (84 targets)")
add_box(s, 7.1, 2.85, 2.4, 0.9, "mγγ adversary", "MLP on μ → 8 SB mγγ bins", fill="FFFFFF", line=RED2)
add_arrow(s, 8.15, 2.55, 0.3, 0.28, MSO_SHAPE.UP_ARROW, GREY)
add_bullets(s, 0.6, 4.0, 6.3, 2.7, [
    (0, [("Data", True), (": HAXAD pseudo-data from the demonstrator pipeline (data pool, unweighted, contrastive-encoder normalization); "
                          "SB = 105–120 ∪ 130–145 GeV, SR = 120–130 GeV", False)]),
    (0, [("Split", True), (": SB k-fold (5, seed): fold 0 test, fold 1 validation, rest training (958k); the SR is never trained on; "
                           "SM Higgs removed in all tests", False)]),
    (0, [("Loss", True), (": masked reconstruction (MSE, BCE for existence and b-tag) + 0.01·KL + 10·(m̂ − mγγ)² "
                          "+ λ·adversary term", False)]),
], size=12)
add_bullets(s, 7.1, 4.0, 5.8, 2.7, [
    (0, [("Photon slots", True), (": pT, η, φ of both photons; the photon-pair mass in the interaction features equals mγγ, "
                                  f"so the encoder reproduces it: m̂ R² {RUN['s1f_adv0p05']['myy_head_r2']:.4f}, "
                                  f"rms {RUN['s1f_adv0p05']['myy_head_rms_gev']:.2f} GeV", False)]),
    (0, [("Adversary", True), (": own Adam optimizer, 20 inner steps per batch on detached μ; the encoder is penalized by "
                               "KL(bin prior ‖ adversary output), switched on after 5 epochs, ramped over 10", False)]),
    (0, [("Training", True), (": 400 epochs, AdamW, batch 16384, bf16; about 3.3 h on one A100", False)]),
], size=12)

# --- 4 study plan --------------------------------------------------------------------------------------------------
s = new_content(); slides.append(s); set_title(s, "Study Plan and Where We Are")
add_steps(s, 0.6, 1.15, 12.1, 0.75,
          ["1 Reconstruction", "1b Fixed mγγ dim", "1c Decorrelate z", "Multi-seed baseline", "2 Weak-supervision test",
           "3 Generation"], ["done", "done", "done", "done", "todo", "todo"])
add_table(s, 0.6, 2.2, 12.1, [2.2, 4.6, 5.3], [
    ["Step", "Question", "Answer so far"],
    ["1  Reconstruction", "Can a ParT CVAE reconstruct the events?",
     f"yes: core R² {ms('step1', 'r2_sb')} (3 seeds); but z carries mγγ and the decoder ignores mγγ"],
    ["1b  Fixed mγγ dim", "Can one latent dim be mγγ itself?",
     f"yes, with photon slots: m̂ R² {RUN['s1f_adv0p05']['myy_head_r2']:.4f}; the decoded photon pair follows the given mγγ"],
    ["1c  Decorrelate z", "Can the rest of z be made mγγ-independent?",
     "yes, with a strong adversary; GRL, DisCo and larger β do not work; price: R² 0.97 → 0.90–0.93"],
    ["Baseline", "Which configuration, and is the gain real?",
     f"λ = 0.05: closure SR/SB {ms('adv0.05', 'auc_srsb')} over 7 seeds vs {ms('step1', 'auc_srsb')} (step 1)"],
    ["2  Weak supervision", "Does z find an injected signal in the SR?", "next (to discuss)"],
    ["3  Generation", "Sample events at a given mγγ", "later"],
], size=10.5, row_h=0.47)
add_text(s, 0.6, 5.65, 12.1, 0.6, "Core R²: median R² over the jet1–4 and fatjet1 targets on the SB test fold. Closure AUCs: "
         "HistGradientBoosting on μ, 3-fold, event-weighted; SB-low vs SB-high on the SB test fold, SR vs SB test.", size=9, italic=True)

# --- 5 step 1 diagnosis -------------------------------------------------------------------------------------------
s = new_content(); slides.append(s); set_title(s, "Step 1: Good Reconstruction, but z Carries mγγ")
add_image(s, "mu_vs_myy.png", 0.5, 1.1, w=7.6)
r0 = RUN["step1_seed0"]
add_bullets(s, 8.3, 1.15, 4.6, 5.4, [
    (0, [("Reconstruction", True), (f": core R² SB {ms('step1', 'r2_sb')}, SR {ms('step1', 'r2_sr')}; "
                                    "b-tag AUC 1.00; MET and high-level ≥ 0.99", False)]),
    (0, [("z knows mγγ", True), (f": closure AUC SB-low vs SB-high {ms('step1', 'auc_lohi')}, SR vs SB "
                                  f"{ms('step1', 'auc_srsb')} (raw inputs: 0.669 / 0.543)", False)]),
    (1, "left: one dim moves by about 0.55 σ across the SB"),
    (0, [("The decoder ignores mγγ", True), (f": core R² {f3(r0['shuffle_r2'][0])} with the true mγγ, "
                                              f"{f3(r0['shuffle_r2'][1])} with mγγ shuffled between events", False)]),
    (1, "everything mγγ-related travels through z, so SR recon = SB recon proved nothing"),
    (0, [("Cause", True), (": the encoder never saw the photons; the mγγ dependence it learns comes from correlated "
                           "kinematics (diphoton pT, jets)", False)]),
], size=12)
add_text(s, 0.6, 4.0, 7.5, 0.5, "Mean of μ_k in mγγ bins (SB test + SR, SM Higgs removed), in units of the dim's standard deviation; "
         "right: the recommended baseline, same scale.", size=9, italic=True)
st1 = [RUN[t] for t in GRP["step1"]["tags"]]
rows = [["Core R² (step 1)", "seed 0", "seed 1", "seed 2"],
        ["SB test", *[f3(r["r2_sb"]) for r in st1]], ["SR", *[f3(r["r2_sr"]) for r in st1]]]
for k in ["XHH_500", "HVT_VcXjjHyy_2000_300", "WN_HyyN_300", "TT_tZNtHyyN_1000_205_60"]:
    rows.append([f"signal {k} (never trained on)", *[f3(r["r2_signal"][k]) for r in st1]])
add_table(s, 0.6, 4.55, 7.5, [3.9, 1.2, 1.2, 1.2], rows, size=9.5, row_h=0.29)

# --- 6 step 1b/1c method -------------------------------------------------------------------------------------------
s = new_content(); slides.append(s); set_title(s, "Steps 1b–1c: Fix mγγ in the Latent, Remove It from z")
d = {t: RUN[t] for t in ["ce_mlp_s0", "ce_ph_mlp_s0", "s1c_b0p1", "s1d_adv0p3", "s1c_adv1", "s1c_adv10", "s1d_nophot_adv1"]}
add_table(s, 0.6, 1.2, 7.6, [3.6, 1.0, 1.0, 1.0], [
    ["Configuration (150 epochs, seed 0)", "AUC lo/hi", "AUC SR/SB", "core R²"],
    ["no photons, no adversary (step-1 recipe)", f3(d["ce_mlp_s0"]["auc_lohi"]), f3(d["ce_mlp_s0"]["auc_srsb"]), f3(d["ce_mlp_s0"]["r2_sb"])],
    ["+ photon slots + m̂ head", f3(d["ce_ph_mlp_s0"]["auc_lohi"]), f3(d["ce_ph_mlp_s0"]["auc_srsb"]), f3(d["ce_ph_mlp_s0"]["r2_sb"])],
    ["  + β = 0.1 (stronger KL)", f3(d["s1c_b0p1"]["auc_lohi"]), f3(d["s1c_b0p1"]["auc_srsb"]), f3(d["s1c_b0p1"]["r2_sb"])],
    ["  + adversary λ = 0.3", f3(d["s1d_adv0p3"]["auc_lohi"]), f3(d["s1d_adv0p3"]["auc_srsb"]), f3(d["s1d_adv0p3"]["r2_sb"])],
    ["  + adversary λ = 1", f3(d["s1c_adv1"]["auc_lohi"]), f3(d["s1c_adv1"]["auc_srsb"]), f3(d["s1c_adv1"]["r2_sb"])],
    ["  + adversary λ = 10", f3(d["s1c_adv10"]["auc_lohi"]), f3(d["s1c_adv10"]["auc_srsb"]), f3(d["s1c_adv10"]["r2_sb"])],
    ["no photons + adversary λ = 1", f3(d["s1d_nophot_adv1"]["auc_lohi"]), f3(d["s1d_nophot_adv1"]["auc_srsb"]), f3(d["s1d_nophot_adv1"]["r2_sb"])],
], size=11, row_h=0.5)
add_text(s, 0.6, 5.3, 7.6, 0.5, "Adversary: 20 inner steps, hidden 256. 150-epoch R² is lower than at 400 epochs for every configuration.",
         size=9, italic=True)
add_bullets(s, 8.5, 1.15, 4.4, 5.4, [
    (0, [("Photon slots make mγγ explicit", True), (": the encoder reproduces it, but without a counter-force it also spreads it "
                                                     "into z (AUC 0.64 → 0.86)", False)]),
    (0, [("A larger β does not remove it", True), (": the KL term compresses z but keeps mγγ (AUC 0.90)", False)]),
    (0, [("The adversary needs to be strong", True), (": with one-step gradient reversal the encoder dodges it (adversary at chance "
                                                      "while a BDT still separates); own optimizer + 20 inner steps work", False)]),
    (0, [("λ sets the trade-off", True), (": λ ≥ 1 reaches the closure floor but destroys the reconstruction; small λ over "
                                          "400 epochs is the useful regime", False)]),
    (0, [("Without photons the decoder ignores mγγ", True), (" even with the adversary, so this branch was dropped", False)]),
], size=12)

# --- 7 multi-seed baseline -----------------------------------------------------------------------------------------
s = new_content(); slides.append(s); set_title(s, "Multi-Seed Baseline: λ = 0.05 Gives Stable Closure")
add_image(s, "tradeoff.png", 0.5, 1.1, w=8.0)
add_table(s, 0.6, 4.4, 7.9, [2.6, 0.6, 1.4, 1.4, 1.4], [
    ["Configuration (400 epochs)", "seeds", "AUC SR/SB", "AUC lo/hi", "core R² SB"],
    ["Step 1 (no photons, no adversary)", "3", ms("step1", "auc_srsb"), ms("step1", "auc_lohi"), ms("step1", "r2_sb")],
    ["photons + adversary λ = 0.03", "7", ms("adv0.03", "auc_srsb"), ms("adv0.03", "auc_lohi"), ms("adv0.03", "r2_sb")],
    ["photons + adversary λ = 0.05", "7", ms("adv0.05", "auc_srsb"), ms("adv0.05", "auc_lohi"), ms("adv0.05", "r2_sb")],
], size=10, row_h=0.34, bold_rows=(3,))
add_text(s, 0.6, 5.85, 7.9, 0.4, "Mean ± standard deviation over seeds (the seed also changes the k-fold split). The closure estimator "
         "itself repeats to ±0.001 (four repetitions).", size=9, italic=True)
g5, g3 = GRP["adv0.05"], GRP["adv0.03"]
add_bullets(s, 8.8, 1.15, 4.1, 5.4, [
    (0, [("λ = 0.05: closure is seed-independent", True), (f": SR/SB {g5['auc_srsb']['min']:.3f}–{g5['auc_srsb']['max']:.3f} "
                                                          f"for all 7 seeds, 0.017 below step 1 (≈ 4× its seed spread)", False)]),
    (0, [("λ = 0.03 is bimodal", True), (f": SR/SB {g3['auc_srsb']['min']:.3f}–{g3['auc_srsb']['max']:.3f}; on average no better than step 1", False)]),
    (0, [("mγγ leakage within the SB is cut", True), (f": lo/hi {ms('step1', 'auc_lohi')} → {ms('adv0.05', 'auc_lohi')}", False)]),
    (0, [("Price: reconstruction", True), (f": R² {g5['r2_sb']['min']:.2f}–{g5['r2_sb']['max']:.2f}, seed-dependent; the worst seeds lose "
                                           "R² without gaining closure", False)]),
    (0, [("The tuning scan is a plateau", True), (" (open circles): latent 32, KL 0.003, stronger or later adversary all fall "
                                                  "inside the seed spread", False)]),
], size=11)

# --- 8 baseline config and the decoder -----------------------------------------------------------------------------
b = RUN["s1f_adv0p05"]
s = new_content(); slides.append(s); set_title(s, "Recommended Baseline and What Its Decoder Does")
add_table(s, 0.6, 1.2, 5.6, [2.2, 3.4], [
    ["Knob", "Baseline value"],
    ["Inputs", "11 object slots + 2 photon slots"],
    ["Latent", "z 16 dims + m̂γγ head (weight 10)"],
    ["Decoder condition", "MLP(mγγ) 1 → 32, concatenated with z"],
    ["KL weight", "0.01"],
    ["Adversary", "λ = 0.05, 8 bins, hidden 256, 20 steps, lr 1e-3"],
    ["Training", "400 epochs; pick the seed by SB validation loss"],
], size=11, row_h=0.42)
add_text(s, 0.6, 4.3, 5.6, 1.2, f"Seed choice: across the 7 λ = 0.05 seeds the validation reconstruction loss ranks the test R² "
         f"with Spearman {g5['spearman_valrecon_r2']:.2f}; it picks seed 0 ({b['tag']}): AUC SR/SB {f3(b['auc_srsb'])}, "
         f"lo/hi {f3(b['auc_lohi'])}, core R² {f3(b['r2_sb'])} (SB) / {f3(b['r2_sr'])} (SR).", size=11)
add_image(s, "myy_response.png", 6.4, 1.15, w=6.5)
add_bullets(s, 6.4, 3.2, 6.5, 3.4, [
    (0, [("Decoded photons carry the given mγγ", True), (f": corr(decoded, given) {b['photon_myy_corr_given']:.3f} when mγγ is "
                                                          f"shuffled, corr with the event's true mγγ {b['photon_myy_corr_true_when_shuffled']:+.3f}; "
                                                          f"rms {b['photon_myy_rms_gev']:.1f} GeV", False)]),
    (0, [("Kinematics follow partly", True), (": at fixed z the decoder shifts diphoton pT and MET with mγγ, in the data's direction", False)]),
    (0, [("Multiplicities do not", True), (": n_jets stays flat (data rises), and the jet-1 pT slope is not in the data", False)]),
    (1, "irrelevant for the z-based classifier; matters only if decoded events are used (step 3)"),
], size=11)
add_text(s, 6.4, 2.95, 6.5, 0.3, "Mean shift vs mγγ at fixed z (lines) and in the data (dots); SB test events.", size=9, italic=True)

# --- 9 status and next ---------------------------------------------------------------------------------------------
s = new_content(); slides.append(s); set_title(s, "Status and Next Steps (for Discussion)")
add_table(s, 0.6, 1.2, 7.8, [2.3, 1.1, 4.4], [
    ["Item", "Status", "Content"],
    ["Step 1 reconstruction", "done", "ParT encoder + slot-query decoder, R² 0.965 (3 seeds)"],
    ["Steps 1b–1c", "done", "photon slots, m̂ head, multi-step adversary"],
    ["Multi-seed baseline", "done", "λ = 0.05: closure SR/SB 0.518 ± 0.002, R² 0.90 ± 0.03 (7 seeds)"],
    ["Signal information in z", "open", "supervised signal vs SR background saturates (AUC ≥ 0.994 for every model and raw inputs)"],
    ["Weak-supervision test", "next", "injected signal in the SR, classifier on z (SR vs SB template)"],
    ["Generation", "later", "p(z) sampling + decoder at a given mγγ"],
], size=11, row_h=0.6)
add_bullets(s, 8.7, 1.15, 4.2, 5.4, [
    (0, [("A  Weak-supervision injection test", True), (" (proposed next): signal at several strengths in the SR; CWoLa-style "
                                                        "classifier on z; significance improvement vs strength, step 1 vs baseline vs raw inputs", False)]),
    (0, [("B  Recover reconstruction at λ = 0.05", True), (": longer training or a λ schedule; worth it only if A shows that R² limits "
                                                          "the sensitivity", False)]),
    (0, [("C  Residual closure 0.518", True), (": measure the fake excess it causes in the weak-supervision test with no "
                                               "injected signal (the null of A)", False)]),
    (0, [("D  Generation (step 3)", True), (": only after A; needs the multiplicity dependence", False)]),
], size=12)

# --- backups -------------------------------------------------------------------------------------------------------
backups = []

# B1 per-seed table
s = new_content(); backups.append(s); set_title(s, "Backup: Every Seed of the Baseline Study")
rows = [["Run", "Config", "seed", "AUC SR/SB", "AUC lo/hi", "R² SB", "R² SR", "γγ rms [GeV]"]]
for g, lab in (("step1", "step 1"), ("adv0.03", "λ 0.03"), ("adv0.05", "λ 0.05")):
    for t in GRP[g]["tags"]:
        r = RUN[t]
        rows.append([t, lab, r["config"]["seed"], f3(r["auc_srsb"]), f3(r["auc_lohi"]), f3(r["r2_sb"]), f3(r["r2_sr"]),
                     f"{r['photon_myy_rms_gev']:.2f}" if "photon_myy_rms_gev" in r else "–"])
add_table(s, 0.6, 1.15, 9.0, [2.2, 1.0, 0.6, 1.1, 1.1, 0.9, 0.9, 1.2], rows, size=9, row_h=0.285)
add_bullets(s, 9.9, 1.15, 3.0, 5.4, [
    (0, "400 epochs each; the seed also sets the k-fold split"),
    (0, "Runs s1e_*, s1f_* are from the tuning scan, base_* from the multi-seed study; same configuration"),
    (0, "γγ rms: mγγ of the decoded photons vs the true mγγ, SB test"),
    (0, f"Spearman(val loss, test R²): λ 0.05 {GRP['adv0.05']['spearman_valrecon_r2']:.2f}, λ 0.03 "
        f"{GRP['adv0.03']['spearman_valrecon_r2']:.2f}, step 1 {GRP['step1']['spearman_valrecon_r2']:.2f}"),
], size=11)

# B2 tuning scan
s = new_content(); backups.append(s); set_title(s, "Backup: Tuning Scan at 400 Epochs (Seed 0)")
scan = [("λ = 0.03 (reference, 7 seeds)", None), ("λ = 0.01", "s1f_adv0p01"), ("λ = 0.03, latent 32", "s1f_lat32"),
        ("λ = 0.03, KL 0.003", "s1f_kl0p003"), ("λ = 0.03, 40 adversary steps", "s1f_stradv"),
        ("λ = 0.03, decoder 256 × 4 layers", "s1f_bigdec"), ("λ = 0.1, adversary from epoch 100", "s1f_adv0p1_late"),
        ("λ = 0.1", "s1e_ph_adv0p1"), ("λ = 0.3", "s1e_ph_adv0p3")]
rows = [["Configuration", "AUC SR/SB", "AUC lo/hi", "R² SB", "R² SR"]]
for lab, t in scan:
    if t is None:
        rows.append([lab, ms("adv0.03", "auc_srsb"), ms("adv0.03", "auc_lohi"), ms("adv0.03", "r2_sb"), ms("adv0.03", "r2_sr")])
    else:
        r = RUN[t]
        rows.append([lab, f3(r["auc_srsb"]), f3(r["auc_lohi"]), f3(r["r2_sb"]), f3(r["r2_sr"])])
add_table(s, 0.6, 1.2, 8.6, [3.2, 1.4, 1.4, 1.3, 1.3], rows, size=11, row_h=0.48)
add_bullets(s, 9.5, 1.15, 3.4, 5.4, [
    (0, "Everything except the decoder change lies inside the λ = 0.03 seed spread"),
    (0, "λ ≥ 0.1 gives the lowest closure AUCs, but R² falls to 0.86–0.91 on one seed; λ = 0.05 is the compromise "
        "that was then run on 7 seeds"),
    (0, "A larger decoder is worse (0.83): the loss is not decoder-capacity-limited"),
], size=12)

# B3 rejected options + COND_EMBED
s = new_content(); backups.append(s); set_title(s, "Backup: Options Tried and Rejected")
add_table(s, 0.6, 1.2, 12.2, [3.0, 5.6, 3.6], [
    ["Option", "Result", "Decision"],
    ["DisCo: λ·dCor²(μ, mγγ)", "unstable at λ ≥ 300 (KL explodes); with warm-up stable, closure 0.562 → 0.551 for 2× worse recon",
     "dropped"],
    ["mγγ as encoder input", "mγγ leaks into z: SB-low vs SB-high AUC 0.85–0.93", "dropped"],
    ["Learned-variance decoder", "KL collapses to about 4 nats", "dropped"],
    ["β = 0.1 (step 1)", "latent starved (about 11 nats), b-tag never learned", "β = 0.01"],
    ["Gradient-reversal adversary (one step)", "encoder dodges it: adversary at chance while a BDT still separates the bins",
     "own optimizer, 20 steps"],
    ["Direct mγγ concatenation (no MLP)", f"same mγγ behaviour; recon {f3(RUN['ce_ph_none_s0']['r2_sb'])} vs "
     f"{f3(RUN['ce_ph_mlp_s0']['r2_sb'])} (photons), {f3(RUN['ce_none_s0']['r2_sb'])} vs {f3(RUN['ce_mlp_s0']['r2_sb'])} "
     "(no photons), 150 epochs", "keep MLP(1 → 32)"],
    ["No photon slots + adversary", "decoder still ignores mγγ (shuffle test, λ = 1: core R² "
     f"{f3(RUN['s1d_nophot_adv1']['shuffle_r2'][0])} → {f3(RUN['s1d_nophot_adv1']['shuffle_r2'][1])})", "dropped"],
    ["Conditioning on multiplicities", "would require specifying them at generation", "rejected by design"],
], size=10, row_h=0.5)
add_text(s, 0.6, 5.85, 12.2, 0.4, "DisCo, encoder-input, learned-variance and β numbers are from the step-1 development runs "
         "(README.md); the others are in figs/summary.json.", size=9, italic=True)

# B4 transfer + signal
tt0, tt1 = RUN["step1_seed0"]["transfer_test"], RUN["s1e_ph_adv0p03"]["transfer_test"]
sig = ["XHH_500", "HVT_VcXjjHyy_2000_300", "WN_HyyN_300", "TT_tZNtHyyN_1000_205_60"]
s = new_content(); backups.append(s); set_title(s, "Backup: Transfer Test and Supervised Signal Sensitivity")
fr, f0 = tt1["A->B"]["mean_shift_fraction"], tt0["A->B"]["mean_shift_fraction"]
add_table(s, 0.6, 1.2, 6.2, [3.0, 1.6, 1.6], [
    ["Decode z_A at mγγ from B (A→B)", "step 1 (seed 0)", "λ = 0.03 (seed 0)"],
    ["BDT AUC transfer vs recon B", f3(tt0["A->B"]["auc_transfer_vs_target"]), f3(tt1["A->B"]["auc_transfer_vs_target"])],
    ["BDT AUC recon A vs recon B", f3(tt0["A->B"]["auc_no_correction"]), f3(tt1["A->B"]["auc_no_correction"])],
    ["shift reproduced: diphoton pT / MET", f"{f2(f0['diphoton_pt'])} / {f2(f0['met_pt'])}", f"{f2(fr['diphoton_pt'])} / {f2(fr['met_pt'])}"],
    ["shift reproduced: jet1 / jet2 pT", f"{f2(f0['jet1_pt'])} / {f2(f0['jet2_pt'])}", f"{f2(fr['jet1_pt'])} / {f2(fr['jet2_pt'])}"],
    ["shift reproduced: n_jets / jet1 existence", f"{f2(f0['n_jets'])} / {f2(f0['jet1_exist'])}", f"{f2(fr['n_jets'])} / {f2(fr['jet1_exist'])}"],
], size=10, row_h=0.36)
add_text(s, 0.6, 3.5, 6.2, 0.7, "A = 105–120 GeV, B = 130–145 GeV (SB test). Shift reproduced = [mean(transfer) − mean(recon A)] / "
         "[mean(recon B) − mean(recon A)]. AUC 0.5 = perfect transfer. Recon A vs B is easy because the decoded photons carry mγγ.", size=9, italic=True)
rows = [["Signal vs SR background, BDT AUC", "raw inputs", "step 1", "λ = 0.03"]]
for k in sig:
    rows.append([k, f3(RUN["step1_seed0"]["signal_sensitivity_raw"][k]["auc"]), f3(RUN["step1_seed0"]["signal_sensitivity"][k]["auc"]),
                 f3(RUN["s1e_ph_adv0p03"]["signal_sensitivity"][k]["auc"])])
add_table(s, 7.1, 1.2, 5.8, [2.8, 1.0, 1.0, 1.0], rows, size=10, row_h=0.36)
add_bullets(s, 7.1, 3.2, 5.8, 3.3, [
    (0, "Fully supervised, 100k events per class: every representation separates every signal, so this test cannot rank models"),
    (0, "Weak supervision at realistic signal fractions is the informative version (proposal A)"),
], size=11)
add_bullets(s, 0.6, 4.25, 6.2, 2.0, [
    (0, "The decoder moves kinematics with mγγ but not the existence logits, so decoded multiplicities do not change"),
    (0, "Not needed for the z-based classifier; relevant again for generation"),
], size=11)

# B5 closure definition
s = new_content(); backups.append(s); set_title(s, "Backup: Closure Test Definition and Its Noise")
add_bullets(s, 0.6, 1.15, 7.4, 5.4, [
    (0, [("Inputs", True), (": μ (16 dims) of events never used in training: the SB test fold and the SR; SM Higgs removed", False)]),
    (0, [("Classifier", True), (": HistGradientBoosting (300 iterations, early stopping), 3-fold, at most 300k events per class, "
                                "classes balanced in the fit, physical weights in the AUC", False)]),
    (0, [("SB-low vs SB-high", True), (": 105–120 vs 130–145 GeV on the SB test fold; the strongest probe of mγγ in z", False)]),
    (0, [("SR vs SB", True), (": SR vs SB test; the quantity that matters for the SR template", False)]),
    (0, [("Estimator noise", True), (": four repetitions with different subsamples give ±0.001 or less; the 0.03 seed spread of "
                                     "λ = 0.03 is model variation", False)]),
], size=12)
add_table(s, 8.3, 1.2, 4.6, [2.0, 1.3, 1.3], [
    ["Run (4 repetitions)", "lo/hi", "SR/SB"],
    ["step1_seed0", "0.6550 ± 0.0003", "0.5382 ± 0.0003"],
    ["step1_seed1", "0.6526 ± 0.0002", "0.5366 ± 0.0004"],
    ["s1e_ph_adv0p03", "0.5709 ± 0.0007", "0.5223 ± 0.0007"],
    ["s1f_adv0p03_s1", "0.5647 ± 0.0008", "0.5472 ± 0.0014"],
], size=9.5, row_h=0.34)
add_text(s, 8.3, 3.05, 4.6, 0.4, "From store/closure_noise.py (logs/closure_noise.log).", size=9, italic=True)

# --- order: title, content, backup divider, backups; drop every other template slide ---------------------------------
want = [title_s] + slides + [backup_s] + backups
lst = prs.slides._sldIdLst
by_id = {int(el.get("id")): el for el in lst}
order = [by_id[sl.slide_id] for sl in want]
for el in list(lst):
    if el not in order:
        prs.part.drop_rel(el.get(qn("r:id")))
    lst.remove(el)
for el in order:
    lst.append(el)
prs.save(OUT)
print("saved", OUT, "with", len(want), "slides")
